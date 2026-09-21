import calendar
import io
from datetime import datetime, timedelta

from flask import Blueprint, flash, redirect, render_template, request, send_file, url_for
from sqlalchemy import func

from app import db
from app.models import Menu, Penjualan
from app.routes.helpers import get_current_user, login_required, pemilik_only

bp = Blueprint("penjualan", __name__)


def _parse_periode(args):
    """Kembalikan (mode, start, end, label, tanggal_str, bulan_str).

    mode: harian | bulanan | semua. start/end naive lokal agar cocok
    dengan <input type=date/month> dan data seed yang naive.
    """
    mode = (args.get("mode") or "bulanan").strip().lower()
    if mode not in ("harian", "bulanan", "semua"):
        mode = "bulanan"
    now = datetime.now()
    tanggal_str = (args.get("tanggal") or now.strftime("%Y-%m-%d")).strip()
    bulan_str = (args.get("bulan") or now.strftime("%Y-%m")).strip()

    start = end = None
    label = "Semua waktu"
    if mode == "harian":
        try:
            d = datetime.strptime(tanggal_str, "%Y-%m-%d")
        except ValueError:
            d = now
            tanggal_str = now.strftime("%Y-%m-%d")
        start = datetime(d.year, d.month, d.day)
        end = start + timedelta(days=1)
        label = start.strftime("%d %b %Y")
    elif mode == "bulanan":
        try:
            d = datetime.strptime(bulan_str, "%Y-%m")
        except ValueError:
            d = now
            bulan_str = now.strftime("%Y-%m")
        start = datetime(d.year, d.month, 1)
        last_day = calendar.monthrange(d.year, d.month)[1]
        end = datetime(d.year, d.month, last_day) + timedelta(days=1)
        label = start.strftime("%b %Y")
    return mode, start, end, label, tanggal_str, bulan_str


def _filtered_query(start, end):
    q = Penjualan.query
    if start and end:
        q = q.filter(Penjualan.waktu >= start, Penjualan.waktu < end)
    return q


def _totals(q):
    omzet = q.with_entities(
        func.coalesce(func.sum(Penjualan.jumlah * Penjualan.harga_satuan), 0)
    ).scalar() or 0
    cup = q.with_entities(func.coalesce(func.sum(Penjualan.jumlah), 0)).scalar() or 0
    count = q.with_entities(func.count(Penjualan.id)).scalar() or 0
    return omzet, cup, count


@bp.route("/penjualan")
@login_required
def daftar():
    mode, start, end, label, tanggal_str, bulan_str = _parse_periode(request.args)
    base = _filtered_query(start, end)
    total_omzet, total_cup, total_trx = _totals(base)
    items = base.order_by(Penjualan.waktu.desc()).limit(200).all()
    menus = Menu.query.filter_by(aktif=True).order_by(Menu.nama).all()
    menus_all = Menu.query.order_by(Menu.nama).all()
    export_url = url_for(
        "penjualan.export", mode=mode, tanggal=tanggal_str, bulan=bulan_str
    )
    return render_template(
        "penjualan.html",
        items=items,
        total_omzet=total_omzet,
        total_cup=total_cup,
        total_trx=total_trx,
        periode_label=label,
        mode=mode,
        tanggal_str=tanggal_str,
        bulan_str=bulan_str,
        menus=menus,
        menus_all=menus_all,
        export_url=export_url,
    )


@bp.route("/penjualan/tambah", methods=["POST"])
@login_required
def tambah():
    nama_menu = request.form.get("nama_menu", "").strip()
    mode = request.form.get("mode", "bulanan")
    tanggal_str = request.form.get("tanggal", "")
    bulan_str = request.form.get("bulan", "")
    try:
        jumlah = int(request.form.get("jumlah", 1))
    except (TypeError, ValueError):
        flash("Jumlah tidak valid.", "danger")
        return redirect(
            url_for("penjualan.daftar", mode=mode, tanggal=tanggal_str, bulan=bulan_str)
        )
    if not nama_menu or jumlah <= 0:
        flash("Nama menu dan jumlah wajib diisi.", "danger")
        return redirect(
            url_for("penjualan.daftar", mode=mode, tanggal=tanggal_str, bulan=bulan_str)
        )
    menu = Menu.query.filter(
        func.lower(Menu.nama) == nama_menu.lower(), Menu.aktif == True  # noqa: E712
    ).first()
    if not menu:
        flash("Pilih menu dari daftar yang tersedia.", "danger")
        return redirect(
            url_for("penjualan.daftar", mode=mode, tanggal=tanggal_str, bulan=bulan_str)
        )
    # Harga dikunci dari master agar kasir tidak salah input/tipu.
    # Waktu pakai jam lokal (naive) agar cocok dengan filter harian/bulanan
    # dan <input type=date> kasir. Jangan pakai UTC: transaksi jam 00-07
    # WIB bakal tercatat mundur sehari dan "hilang" dari filter hari ini.
    harga = float(menu.harga_default or 0)
    user = get_current_user()
    db.session.add(
        Penjualan(
            nama_menu=menu.nama,
            jumlah=jumlah,
            harga_satuan=harga,
            waktu=datetime.now(),
            dicatat_oleh=user.id,
        )
    )
    db.session.commit()
    flash(f"Penjualan '{menu.nama}' x{jumlah} dicatat.", "success")
    return redirect(
        url_for("penjualan.daftar", mode=mode, tanggal=tanggal_str, bulan=bulan_str)
    )


@bp.route("/penjualan/<int:pid>/hapus", methods=["POST"])
@login_required
@pemilik_only
def hapus(pid):
    p = Penjualan.query.get_or_404(pid)
    mode = request.form.get("mode", "bulanan")
    tanggal_str = request.form.get("tanggal", "")
    bulan_str = request.form.get("bulan", "")
    nama = p.nama_menu
    db.session.delete(p)
    db.session.commit()
    flash(f"Transaksi '{nama}' dihapus dari riwayat.", "info")
    return redirect(
        url_for("penjualan.daftar", mode=mode, tanggal=tanggal_str, bulan=bulan_str)
    )


@bp.route("/penjualan/export")
@login_required
def export():
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    mode, start, end, label, _, _ = _parse_periode(request.args)
    rows = (
        _filtered_query(start, end)
        .order_by(Penjualan.waktu.asc())
        .limit(5000)
        .all()
    )
    total_omzet, total_cup, total_trx = _totals(_filtered_query(start, end))

    wb = Workbook()
    ws = wb.active
    ws.title = "Penjualan"
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    header_font = Font(bold=True, color="FFFFFF", size=10)
    header_fill = PatternFill(start_color="4B2E1E", end_color="4B2E1E", fill_type="solid")
    title_font = Font(bold=True, size=13, color="22140E")
    thin = Side(style="thin", color="E8E0D6")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    ws.merge_cells("A1:F1")
    ws["A1"] = f"Rekap Penjualan — {label}"
    ws["A1"].font = title_font
    ws["A2"] = f"Diekspor: {datetime.now().strftime('%d %b %Y %H:%M')} | Transaksi: {total_trx} | Cup: {total_cup} | Omzet: Rp {float(total_omzet or 0):,.0f}"
    ws["A2"].font = Font(size=9, color="6E645E")
    ws.merge_cells("A2:F2")

    headers = ["Waktu", "Menu", "Jumlah (cup)", "Harga Satuan (Rp)", "Subtotal (Rp)", "Dicatat Oleh"]
    for col, h in enumerate(headers, start=1):
        cell = ws.cell(row=4, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border
    ws.row_dimensions[4].height = 22

    # Pencatat: nama user bila tersedia
    from app.models import Pengguna

    pencatat = {u.id: u.nama for u in Pengguna.query.all()}
    r = 5
    for p in rows:
        waktu = p.waktu.strftime("%Y-%m-%d %H:%M") if p.waktu else "-"
        subtotal = float(p.jumlah or 0) * float(p.harga_satuan or 0)
        vals = [
            waktu,
            p.nama_menu,
            int(p.jumlah or 0),
            float(p.harga_satuan or 0),
            subtotal,
            pencatat.get(p.dicatat_oleh, "-"),
        ]
        for col, v in enumerate(vals, start=1):
            cell = ws.cell(row=r, column=col, value=v)
            cell.font = Font(size=10)
            cell.border = border
            if col in (3, 4, 5):
                cell.number_format = "#,##0"
                cell.alignment = Alignment(horizontal="right")
            elif col == 1:
                cell.alignment = Alignment(horizontal="center")
        r += 1

    # Baris total
    for col, v in enumerate(
        ["TOTAL", "", int(total_cup or 0), "", float(total_omzet or 0), f"{total_trx} transaksi"],
        start=1,
    ):
        cell = ws.cell(row=r, column=col, value=v)
        cell.font = Font(bold=True, size=10)
        cell.fill = PatternFill(start_color="FBF2E8", end_color="FBF2E8", fill_type="solid")
        cell.border = border
        if col in (3, 5):
            cell.number_format = "#,##0"
            cell.alignment = Alignment(horizontal="right")

    widths = [17, 26, 14, 18, 18, 18]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(row=4, column=i).column_letter].width = w
    ws.freeze_panes = "A5"
    ws.auto_filter.ref = f"A4:F{r - 1 if r > 5 else 5}"

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    fname = f"penjualan_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"
    return send_file(
        buf,
        as_attachment=True,
        download_name=fname,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
